"""Turn an uploaded file (xlsx, csv, docx, pdf, txt) into plain text for the model."""
import csv
import io
import re
import zipfile

MAX_CHARS = 60_000


class ParseError(Exception):
    pass


def to_text(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext in ("xlsx", "xlsm"):
        text = _xlsx(data)
    elif ext == "csv":
        text = _csv(data)
    elif ext == "docx":
        text = _docx(data)
    elif ext == "pdf":
        text = _pdf(data)
    elif ext in ("txt", "md", "json"):
        text = _decode(data)
    else:
        raise ParseError(f"Unsupported file type: .{ext}. Use xlsx, csv, docx, pdf or txt.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ParseError("The file has no readable text.")
    return text[:MAX_CHARS]


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _csv(data: bytes) -> str:
    rows = csv.reader(io.StringIO(_decode(data)))
    return "\n".join(" | ".join(c.strip() for c in row) for row in rows if any(c.strip() for c in row))


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ParseError(f"Could not open the Excel file: {e}")
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                out.append(" | ".join(cells).rstrip(" |"))
    return "\n".join(out)


def _docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    except Exception as e:
        raise ParseError(f"Could not open the Word file: {e}")
    xml = re.sub(r"</w:tc>", " | ", xml)  # table cells
    xml = re.sub(r"</w:p>|</w:tr>", "\n", xml)  # paragraphs and table rows
    text = re.sub(r"<[^>]+>", "", xml)
    for a, b in (("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"'), ("&apos;", "'")):
        text = text.replace(a, b)
    return text


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages[:50])
    except Exception as e:
        raise ParseError(f"Could not read the PDF: {e}")
